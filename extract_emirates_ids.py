"""Extract Emirates ID fields, portrait photos, and valid full ID-card images."""

from __future__ import annotations

import argparse
import io
import re
from dataclasses import dataclass
from pathlib import Path

import cv2
import numpy as np
import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageOps
from rapidocr_onnxruntime import RapidOCR


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}
ID_PATTERN = re.compile(r"(?<!\d)784\D{0,3}(\d{4})\D{0,3}(\d{7})\D{0,3}(\d)(?!\d)")
NAME_PATTERN = re.compile(r"\bname\s*[:：.]?\s*(.*)", re.IGNORECASE)
FACE_MODEL = Path(__file__).parent / "data/haarcascades/haarcascade_frontalface_default.xml"


@dataclass
class TextLine:
    text: str
    confidence: float
    x: float
    y: float
    height: float


@dataclass
class Result:
    source: str
    page: int
    name: str = ""
    id_number: str = ""
    photo: str = ""
    id_image: str = ""
    status: str = "review"
    details: str = ""


def images_in_file(path: Path):
    """Yield page images, using the original PDF image when possible."""
    if path.suffix.lower() != ".pdf":
        with Image.open(path) as image:
            yield 1, ImageOps.exif_transpose(image).convert("RGB")
        return
    with pymupdf.open(path) as document:
        for page_number, page in enumerate(document, 1):
            images = page.get_images(full=True)
            if len(images) == 1:
                xref = images[0][0]
                rects = page.get_image_rects(xref)
                if rects and rects[0].get_area() >= page.rect.get_area() * 0.25:
                    data = document.extract_image(xref)["image"]
                    with Image.open(io.BytesIO(data)) as image:
                        yield page_number, ImageOps.exif_transpose(image).convert("RGB")
                    continue
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(3, 3), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            pixels = np.asarray(image)
            nonwhite = np.any(pixels < 240, axis=2)
            if nonwhite.any():
                ys, xs = np.where(nonwhite)
                image = image.crop((int(xs.min()), int(ys.min()), int(xs.max()) + 1, int(ys.max()) + 1))
            yield page_number, image


def read_lines(image: Image.Image, engine: RapidOCR) -> list[TextLine]:
    # Keep text in its physical orientation: auto-flipping individual OCR crops
    # would make an upside-down card look readable without rotating the card.
    raw, _ = engine(np.asarray(image), use_cls=False)
    lines = []
    for box, text, confidence in raw or []:
        xs = [point[0] for point in box]
        ys = [point[1] for point in box]
        lines.append(TextLine(text, float(confidence), min(xs), min(ys), max(ys) - min(ys)))
    return lines


def english_name(value: str) -> str:
    value = re.split(r"[^\x00-\x7f]|\b(?:Date of Birth|Nationality|ID Number)\b", value)[0]
    value = re.sub(r"[^A-Za-z .'-]", " ", value)
    return re.sub(r"\s+", " ", value).strip(" .'-")


def name_from_lines(label: TextLine, lines: list[TextLine]) -> tuple[str, float]:
    """Read the English name block, including rows above a centered label."""
    match = NAME_PATTERN.search(label.text)
    inline = english_name(match[1]) if match else ""
    height = max(label.height, 1.0)
    # For a two-row name, the first row starts within one label height.
    # The Arabic name sits above it and can occasionally be misread as Latin.
    top = label.y - height
    bottom = label.y + 3 * height
    field_pattern = re.compile(
        r"\b(?:ID\s*Number|Date\s*of\s*Birth|Nationality|Issuing\s*Date|"
        r"Expiry\s*Date|Sex|Signature)\b", re.IGNORECASE
    )
    # Stop at adjacent fields even when their values are separate OCR boxes.
    for other in lines:
        if ID_PATTERN.search(other.text) or field_pattern.search(other.text):
            if other.y + other.height <= label.y:
                top = max(top, other.y + other.height)
            elif other.y > label.y:
                bottom = min(bottom, other.y)

    parts = [(label, inline)] if inline else []
    for other in lines:
        if other is label or NAME_PATTERN.search(other.text) or field_pattern.search(other.text):
            continue
        if not (top <= other.y < bottom and other.x > label.x + height):
            continue
        # Arabic text and dates are not English name continuations.
        if not re.fullmatch(r"[A-Za-z .'-]+", other.text.strip()):
            continue
        value = english_name(other.text)
        if value:
            parts.append((other, value))
    if not parts:
        return "", 0.0

    # Group boxes on the same row before sorting horizontally. The label may
    # sit slightly lower than an inline name box due to OCR box geometry.
    parts.sort(key=lambda item: (item[0].y, item[0].x))
    rows = []
    for part in parts:
        if rows and abs(part[0].y - rows[-1][0][0].y) < 0.5 * max(part[0].height, rows[-1][0][0].height):
            rows[-1].append(part)
        else:
            rows.append([part])
    ordered = [part for row in rows for part in sorted(row, key=lambda item: item[0].x)]
    name = " ".join(value for _, value in ordered).title()
    return name, min(part.confidence for part, _ in ordered)


def fields_from_lines(lines: list[TextLine]) -> tuple[str, str, float, float]:
    ids = []
    names = []
    for line in lines:
        match = ID_PATTERN.search(line.text)
        if match:
            ids.append((f"784-{match[1]}-{match[2]}-{match[3]}", line.confidence))
        if NAME_PATTERN.search(line.text):
            name, confidence = name_from_lines(line, lines)
            if name:
                names.append((name, confidence))
    name, name_score = max(names, key=lambda item: item[1], default=("", 0.0))
    id_number, id_score = max(ids, key=lambda item: item[1], default=("", 0.0))
    return name, id_number, name_score, id_score


def extract_fields(image: Image.Image, engine: RapidOCR):
    best = ("", "", 0.0, 0.0, image)
    best_rank = (0, 0, 0, 0.0)
    for angle in (0, 180, 90, 270):
        rotated = image.rotate(angle, expand=True) if angle else image
        # The supplied scans are under 1000 px wide; modest enlargement keeps
        # the English letter shapes legible to the OCR recognizer.
        ocr_image = rotated
        if rotated.width < 1600:
            ocr_image = rotated.resize((round(rotated.width * 1.5), round(rotated.height * 1.5)))
        lines = read_lines(ocr_image, engine)
        name, number, name_score, id_score = fields_from_lines(lines)
        if number:
            # Compare both scales: scans may lose the small Name label,
            # continuation, or spaces between words at either resolution.
            enlarged = rotated.resize((rotated.width * 2, rotated.height * 2))
            larger_lines = read_lines(enlarged, engine)
            # Preserve word spacing from the first pass when both passes read
            # the same letters but the larger pass joins adjacent words.
            for larger in larger_lines:
                for original in lines:
                    if (re.sub(r"\s+", "", larger.text).casefold()
                            == re.sub(r"\s+", "", original.text).casefold()
                            and original.text.count(" ") > larger.text.count(" ")):
                        larger.text = original.text
                        break
            retry = fields_from_lines(larger_lines)
            if retry[0] and (not name or len(retry[0].split()) >= len(name.split())):
                name, name_score = retry[0], retry[2]
            if retry[1] and retry[3] > id_score:
                number, id_score = retry[1], retry[3]
            lines = larger_lines
        upright_labels = sum(bool(re.search(
            r"\b(?:Name|ID\s*Number|Date\s*of\s*Birth|Nationality|"
            r"Resident\s*Identity|UNITED\s*ARAB\s*EMIRATES)\b",
            line.text, re.IGNORECASE,
        )) for line in lines)
        # OCR also normalizes vertical text crops. A card front should be
        # landscape, so a sideways recognition must not end the search early.
        landscape = rotated.width >= rotated.height
        rank = (bool(name) + bool(number), int(landscape), upright_labels, name_score + id_score)
        if rank > best_rank:
            best = (name, number, name_score, id_score, rotated)
            best_rank = rank
        if landscape and name and number and min(name_score, id_score) >= 0.85:
            break
    if not best[0] or not best[1]:
        enlarged = best[4].resize((best[4].width * 2, best[4].height * 2))
        name, number, name_score, id_score = fields_from_lines(read_lines(enlarged, engine))
        if name and (not best[0] or name_score > best[2]):
            best = (name, best[1], name_score, best[3], best[4])
        if number and (not best[1] or id_score > best[3]):
            best = (best[0], number, best[2], id_score, best[4])
    return best


def portrait_crop(image: Image.Image, classifier: cv2.CascadeClassifier):
    pixels = np.asarray(image)
    height, width = pixels.shape[:2]
    scale = min(1.0, 1200 / width)
    small = cv2.resize(pixels, None, fx=scale, fy=scale) if scale < 1 else pixels
    gray = cv2.cvtColor(small, cv2.COLOR_RGB2GRAY)
    faces = classifier.detectMultiScale(gray, scaleFactor=1.05, minNeighbors=5, minSize=(30, 30))
    if len(faces) == 0:
        return None
    x, y, face_width, _ = max(faces, key=lambda face: face[2] * face[3])
    x, y, face_width = (float(value) / scale for value in (x, y, face_width))
    if x + face_width / 2 < width / 2:
        top = 0.22 if y / height < 0.31 else 0.26
        box = (0.05 * width, top * height, 0.31 * width, 0.70 * height)
    else:
        box = (0.68 * width, 0.47 * height, 0.91 * width, height)
    return image.crop(tuple(round(value) for value in box))


def unique_jpg_path(folder: Path, name: str, used: set[str]) -> Path:
    base = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", name).strip(" .")[:120] or "Unknown"
    candidate = folder / f"{base}.jpg"
    counter = 2
    while candidate.name.casefold() in used:
        candidate = folder / f"{base}_{counter}.jpg"
        counter += 1
    used.add(candidate.name.casefold())
    return candidate


def write_workbook(rows: list[Result], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Emirates IDs"
    sheet.append(["Source file", "Page", "English name", "Emirates ID number", "Photo", "ID image", "Status", "Details"])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in rows:
        sheet.append([row.source, row.page, row.name, row.id_number, row.photo, row.id_image, row.status, row.details])
        sheet.cell(sheet.max_row, 4).number_format = "@"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 34, "B": 8, "C": 48, "D": 26, "E": 52, "F": 52, "G": 12, "H": 50}.items():
        sheet.column_dimensions[column].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def process(input_dir: Path, output_dir: Path) -> list[Result]:
    if not input_dir.is_dir():
        raise FileNotFoundError(f"Input folder does not exist: {input_dir}")
    photos_dir = output_dir / "photos"
    id_cards_dir = output_dir / "id_cards"
    photos_dir.mkdir(parents=True, exist_ok=True)
    id_cards_dir.mkdir(parents=True, exist_ok=True)
    engine = RapidOCR()
    classifier = cv2.CascadeClassifier(str(FACE_MODEL))
    if classifier.empty():
        raise RuntimeError(f"Could not load face model: {FACE_MODEL}")

    rows = []
    used_photo_names = set()
    used_id_card_names = set()
    files = sorted(path for path in input_dir.iterdir() if path.is_file() and (path.suffix.lower() in IMAGE_EXTENSIONS or path.suffix.lower() == ".pdf"))
    for path in files:
        try:
            for page_number, image in images_in_file(path):
                row = Result(source=path.name, page=page_number)
                try:
                    row.name, row.id_number, name_score, id_score, oriented = extract_fields(image, engine)

                    # A PDF page is considered a valid Emirates ID page only when
                    # both the English name and Emirates ID number are detected.
                    # Invalid PDF pages are ignored completely: they are not saved
                    # as images and they are not written to the workbook.
                    is_valid_id_page = bool(row.name and row.id_number)
                    if path.suffix.lower() == ".pdf" and not is_valid_id_page:
                        missing = []
                        if not row.name:
                            missing.append("English name")
                        if not row.id_number:
                            missing.append("Emirates ID number")
                        print(
                            f"{path.name} page {page_number}: skipped invalid PDF page"
                            f" | missing {', '.join(missing) or 'required fields'}"
                        )
                        continue

                    issues = []
                    if not row.name:
                        issues.append("English name not found")
                    elif name_score < 0.85:
                        issues.append("Low confidence name")
                    if not row.id_number:
                        issues.append("Emirates ID number not found")
                    elif id_score < 0.85:
                        issues.append("Low confidence ID number")

                    # Save the complete input ID/page as <person name>.jpg whenever
                    # the required ID fields were successfully detected.
                    if is_valid_id_page:
                        id_target = unique_jpg_path(id_cards_dir, row.name, used_id_card_names)
                        oriented.save(id_target, "JPEG", quality=95)
                        row.id_image = str(id_target)

                    photo = portrait_crop(oriented, classifier)
                    if photo is None:
                        issues.append("Portrait not found")
                    elif row.name:
                        target = unique_jpg_path(photos_dir, row.name, used_photo_names)
                        photo.save(target, "JPEG", quality=95)
                        row.photo = str(target)

                    row.status = "ok" if not issues else "review"
                    row.details = "; ".join(issues)
                except Exception as exc:
                    if path.suffix.lower() == ".pdf":
                        print(
                            f"{path.name} page {page_number}: skipped PDF page due to error"
                            f" | {type(exc).__name__}: {exc}"
                        )
                        continue
                    row.status = "error"
                    row.details = f"{type(exc).__name__}: {exc}"
                rows.append(row)
                print(f"{path.name} page {page_number}: {row.status} | {row.name or '?'} | {row.id_number or '?'}")
        except Exception as exc:
            rows.append(Result(source=path.name, page=0, status="error", details=f"{type(exc).__name__}: {exc}"))
            print(f"{path.name}: error | {exc}")

    workbook_path = output_dir / "emirates_ids.xlsx"
    write_workbook(rows, workbook_path)
    print(f"Saved {len(rows)} rows to {workbook_path}")
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=Path("input"))
    parser.add_argument("--output", type=Path, default=Path("output"))
    args = parser.parse_args()
    process(args.input, args.output)


if __name__ == "__main__":
    main()
